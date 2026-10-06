from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
def build_dashboard(holdings_df,qoq_df,output):
 output=Path(output); output.parent.mkdir(parents=True,exist_ok=True); summary=pd.DataFrame({'Metric':['Positions','Total market value (USD)','New','Increased','Decreased','Sold out','Unchanged'],'Value':[len(holdings_df),int(holdings_df.get('value_usd',pd.Series(dtype=float)).sum()),int((qoq_df.get('action',pd.Series(dtype=str))=='NEW').sum()),int((qoq_df.get('action',pd.Series(dtype=str))=='INCREASED').sum()),int((qoq_df.get('action',pd.Series(dtype=str))=='DECREASED').sum()),int((qoq_df.get('action',pd.Series(dtype=str))=='SOLD_OUT').sum()),int((qoq_df.get('action',pd.Series(dtype=str))=='UNCHANGED').sum())]})
 with pd.ExcelWriter(output,engine='openpyxl') as w: summary.to_excel(w,sheet_name='Summary',index=False); holdings_df.to_excel(w,sheet_name='Holdings',index=False); qoq_df.to_excel(w,sheet_name='QoQ Changes',index=False)
 wb=load_workbook(output)
 for ws in wb.worksheets:
  ws.freeze_panes='A2'
  for cell in ws[1]: cell.font=Font(bold=True)
  ws.auto_filter.ref=ws.dimensions
 wb.save(output); return output
